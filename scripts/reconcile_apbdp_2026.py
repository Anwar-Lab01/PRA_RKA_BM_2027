#!/usr/bin/env python3
"""Dry-run and idempotent importer for REKONSILIASI_APBDP_2026.xlsx.

The default mode is read-only. It audits the workbook, matches eligible road
packages against the same hardcoded road source consumed by BudgetMap, and
writes CSV audit reports. Database writes require both --execute and
--approve-dry-run. The importer never deletes rows and never overwrites a
manual spatial link.

This script uses the bundled/runtime ``openpyxl`` package for XLSX reading and
the Supabase REST API through Python's standard library. No repository-local
runtime dependency is required.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover - environment guidance
    raise SystemExit("openpyxl diperlukan untuk membaca workbook XLSX") from exc


REPO_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = Path(r"F:\Document\Download\REKONSILIASI_APBDP_2026.xlsx")
DEFAULT_ROADS = REPO_ROOT / "src" / "lib" / "legacyCoordinates.js"
DEFAULT_OUTPUT = REPO_ROOT / "docs" / "audit"
ACTION_SHEETS = ("F6_STAGE", "UPDATE_EXISTING", "CANCEL", "REPLACE")
TARGET_JENIS = {"PLANNING", "PHYSICAL"}

REPORT_COLUMNS = [
    "source_sheet", "source_row", "target_year", "jenis", "nama_paket_source",
    "sub_source", "effective_pagu", "existing_id", "target_anggaran_id",
    "package_action", "road_match_status", "matched_spasial_ref",
    "matched_spasial_nama", "match_score", "match_method", "second_best_score",
    "notes", "database_status",
]

MANUAL_COLUMNS = [
    "source_sheet", "source_row", "target_year", "jenis", "nama_paket_source",
    "sub_source", "effective_pagu", "package_action", "road_match_status",
    "notes", "top_candidates",
]

WORK_PREFIXES = [
    "pengamanan badan jalan", "peningkatan struktur", "siring pasangan batu",
    "pemeliharaan berkala", "pemeliharaan rutin", "jasa konsultansi",
    "long segment", "rekonstruksi", "rehabilitasi", "pembangunan", "pelebaran",
    "pemeliharaan", "pengamanan", "perencanaan", "pengawasan",
]
STOP_WORDS = {
    "badan", "desa", "ds", "jalan", "jasa", "kec", "kecamatan", "konsultansi",
    "lanjutan", "perencanaan", "pengawasan", "ruas", "tahun",
}


def clean_display(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("�", "-")).strip()


def normalize_text(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).lower()
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = text.replace("�", "-").replace("–", "-").replace("—", "-")
    text = re.sub(r"(?<![a-z0-9])(?:jl|jln)\.?\s*", "jalan ", text)
    text = re.sub(r"(?<![a-z0-9])ds\.?\s*", "desa ", text)
    text = re.sub(r"(?<![a-z0-9])sp\.?\s*(\d*)\b", r"simpang \1", text)
    text = re.sub(r"(?<![a-z0-9])sei\.?\b", "sungai", text)
    text = re.sub(r"\b(?:kec|kecamatan)\.?\s+[a-z0-9][a-z0-9 .'/ -]*$", " ", text)
    text = re.sub(r"\btahun\s*20\d{2}\b", " ", text)

    changed = True
    while changed:
        changed = False
        for prefix in sorted(WORK_PREFIXES, key=len, reverse=True):
            pattern = rf"^{re.escape(prefix)}\s+"
            if re.match(pattern, text, flags=re.IGNORECASE):
                text = re.sub(pattern, "", text, count=1, flags=re.IGNORECASE)
                changed = True
                break

    text = re.sub(r"[+/&]", " ", text)
    text = re.sub(r"\s*-\s*", " - ", text)
    text = re.sub(r"[^a-z0-9\s-]", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    segments = []
    for segment in text.split(" - "):
        tokens = [
            token for token in re.split(r"\s+", segment)
            if token and token not in STOP_WORDS and not re.fullmatch(r"20\d{2}", token)
        ]
        if tokens:
            segments.append(" ".join(tokens))
    return " - ".join(segments).strip()


def normalize_package_name(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).lower().replace("�", "-").replace("–", "-").replace("—", "-")
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def normalize_sub_name(value: Any) -> str:
    text = clean_display(value)
    text = re.sub(r"^\s*\d{4}\s*(?:[-:]\s*)?", "", text)
    text = unicodedata.normalize("NFKD", text).lower()
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def extract_sub_name(value: Any) -> str:
    return normalize_sub_name(value)


def source_tokens(value: Any) -> list[str]:
    return normalize_text(value).replace("-", " ").split()


def levenshtein_similarity(left: str, right: str) -> float:
    if left == right:
        return 1.0
    if not left or not right:
        return 0.0
    previous = list(range(len(right) + 1))
    for index, left_char in enumerate(left, start=1):
        current = [index]
        for right_index, right_char in enumerate(right, start=1):
            cost = 0 if left_char == right_char else 1
            current.append(min(current[-1] + 1, previous[right_index] + 1, previous[right_index - 1] + cost))
        previous = current
    return 1 - previous[-1] / max(len(left), len(right))


def score_road_candidate(source_text: str, road: dict[str, Any]) -> dict[str, Any] | None:
    road_ref = road.get("ref") or road.get("id")
    road_name = road.get("input") or road.get("name") or road.get("match") or road_ref
    source_norm = normalize_text(source_text)
    road_norm = normalize_text(road_name)
    if not road_ref or not source_norm or not road_norm:
        return None

    source_flat = re.sub(r"\s+", " ", source_norm)
    road_flat = re.sub(r"\s+", " ", road_norm)
    source_set = set(source_tokens(source_norm))
    road_set = set(source_tokens(road_norm))
    overlap = len(source_set.intersection(road_set))
    coverage = overlap / len(road_set) if road_set else 0.0
    jaccard = overlap / len(source_set.union(road_set)) if source_set.union(road_set) else 0.0
    source_parts = [part for part in source_norm.split(" - ") if part]
    road_parts = [part for part in road_norm.split(" - ") if part]
    reverse_endpoint = len(source_parts) == len(road_parts) >= 2 and source_parts[::-1] == road_parts

    method = "fuzzy_token_edit"
    reason = "Token overlap dan normalized edit similarity."
    if source_flat == road_flat:
        score, method, reason = 1.0, "exact_normalized", "Nama paket/lokasi dan nama ruas sama setelah normalisasi."
    elif road_flat in source_flat:
        score, method, reason = 0.98, "strong_containment", "Nama ruas termuat lengkap pada nama paket/lokasi."
    elif reverse_endpoint:
        score, method, reason = 0.97, "reverse_endpoint", "Endpoint ruas cocok dalam urutan terbalik."
    else:
        score = coverage * 0.5 + jaccard * 0.3 + levenshtein_similarity(source_flat, road_flat) * 0.2
        if coverage >= 0.8 and overlap >= 2:
            score = max(score, 0.86)

    return {
        "road_ref": str(road_ref), "road_name": clean_display(road_name),
        "score": round(min(1.0, score), 4), "method": method, "reason": reason,
        "normalized_source": source_norm, "normalized_road": road_norm,
    }


def is_multi_road_name(value: str) -> bool:
    text = str(value or "").lower()
    return bool(re.search(r"\s(?:dan|\+)\s", text) or re.search(r"\s/\s", text))


def non_road_reason(value: str) -> str | None:
    text = str(value or "").lower()
    if re.search(r"\bjembatan\b|\bbridge\b", text):
        return "Paket jembatan diarahkan ke layer jembatan, bukan polyline jalan."
    if re.search(r"survey kondisi|penyelidikan tanah|dokumen ukl|appraisal objek|studi kelayakan", text):
        return "Paket bersifat studi/dokumen umum dan tidak memiliki identitas ruas yang aman."
    if re.fullmatch(r"\s*pemeliharaan rutin jalan\s*", text):
        return "Pemeliharaan rutin jalan tidak menyebut ruas tertentu."
    return None


def classify_road_match(source_text: str, roads: list[dict[str, Any]]) -> dict[str, Any]:
    skip_reason = non_road_reason(source_text)
    if skip_reason:
        return {"status": "SKIPPED_NON_ROAD", "candidates": [], "notes": skip_reason}
    candidates = sorted(
        (candidate for candidate in (score_road_candidate(source_text, road) for road in roads) if candidate),
        key=lambda item: item["score"], reverse=True,
    )
    best = candidates[0] if candidates else None
    second = candidates[1] if len(candidates) > 1 else None
    if not best or best["score"] < 0.7:
        return {"status": "UNMATCHED", "candidates": candidates[:5], "notes": "Tidak ada kandidat jalan dengan skor minimal 0.70."}
    if is_multi_road_name(source_text):
        return {"status": "REVIEW", "candidates": candidates[:5], "notes": "Nama paket mengindikasikan lebih dari satu ruas; tidak memilih satu ruas secara otomatis."}
    gap = best["score"] - second["score"] if second else best["score"]
    if best["score"] >= 0.95 and gap >= 0.08:
        status = "EXACT"
    elif best["score"] >= 0.85 and gap >= 0.08:
        status = "FUZZY_AUTO"
    elif best["score"] >= 0.7:
        note = "Kandidat terbaik terlalu dekat dengan kandidat kedua." if second and gap < 0.08 else "Skor berada pada rentang review."
        return {"status": "REVIEW", "candidates": candidates[:5], "notes": note}
    else:
        status = "UNMATCHED"
    return {"status": status, "candidates": candidates[:5], "notes": best["reason"]}


def extract_sub_code(value: Any) -> str:
    match = re.match(r"\s*(\d{4})", clean_display(value))
    return match.group(1) if match else ""


def read_sheet_rows(worksheet, header_row: int, start_row: int) -> list[dict[str, Any]]:
    headers = [clean_display(cell.value).lower() for cell in worksheet[header_row]]
    rows = []
    for row_number, row in enumerate(worksheet.iter_rows(min_row=start_row, values_only=True), start=start_row):
        values = list(row[: len(headers)])
        if not any(value is not None for value in values):
            continue
        rows.append({headers[index]: values[index] if index < len(values) else None for index in range(len(headers))} | {"source_row": row_number})
    return rows


def read_workbook(path: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    main_rows = read_sheet_rows(workbook["REKONSILIASI_APBDP_2026"], 8, 9)
    action_counts: dict[str, Any] = {}
    all_rows_for_sub_labels = list(main_rows)
    for sheet_name in ACTION_SHEETS:
        if sheet_name in workbook.sheetnames:
            action_rows = read_sheet_rows(workbook[sheet_name], 1, 2)
            all_rows_for_sub_labels.extend(action_rows)
            action_counts[sheet_name] = {
                "rows": len(action_rows),
                "jenis": dict(Counter(clean_display(row.get("jenis")) or "(blank)" for row in action_rows)),
                "status": dict(Counter(clean_display(row.get("status")) or "(blank)" for row in action_rows)),
                "tindakan": dict(Counter(clean_display(row.get("tindakan")) or "(blank)" for row in action_rows)),
            }

    sub_labels_by_code: dict[str, str] = {}
    for row in all_rows_for_sub_labels:
        code = extract_sub_code(row.get("sub"))
        label = extract_sub_name(row.get("sub"))
        if code and label and code not in sub_labels_by_code:
            sub_labels_by_code[code] = label

    records = []
    for row in main_rows:
        jenis = clean_display(row.get("jenis")).upper()
        status = clean_display(row.get("status")).upper()
        tindakan = clean_display(row.get("tindakan")).upper()
        if jenis not in {"PLANNING", "PHYSICAL"}:
            continue
        records.append({
            "source_sheet": "REKONSILIASI_APBDP_2026", "source_row": row["source_row"],
            "target_year": 2027 if jenis == "PLANNING" else 2026, "jenis": jenis,
            "nama_paket_source": str(row.get("nama") or ""), "nama_paket": clean_display(row.get("nama")),
            "sub_source": str(row.get("sub") or ""), "sub_code": extract_sub_code(row.get("sub")),
            "sub_name": extract_sub_name(row.get("sub")) or sub_labels_by_code.get(extract_sub_code(row.get("sub")), ""),
            "effective_pagu": float(row.get("pagu setelah perubahan") or 0),
            "existing_id": clean_display(row.get("existing id")) or "", "status": status,
            "tindakan": tindakan, "catatan_source": clean_display(row.get("catatan")),
        })
    return records, {
        "sheet_names": workbook.sheetnames, "main_rows": len(main_rows), "action_counts": action_counts,
        "all_main_jenis": dict(Counter(clean_display(row.get("jenis")) or "(blank)" for row in main_rows)),
        "all_main_status": dict(Counter(clean_display(row.get("status")) or "(blank)" for row in main_rows)),
        "all_main_tindakan": dict(Counter(clean_display(row.get("tindakan")) or "(blank)" for row in main_rows)),
        "sub_labels_by_code": sub_labels_by_code,
    }


def load_road_source(path: Path) -> list[dict[str, Any]]:
    text = path.read_text(encoding="utf-8")
    match = re.search(r"export\s+const\s+legacyCoordinates\s*=\s*(\[.*?\])\s*;", text, flags=re.DOTALL)
    if not match:
        raise ValueError(f"Array legacyCoordinates tidak ditemukan pada {path}")
    return json.loads(match.group(1))


def load_env(path: Path) -> dict[str, str]:
    values: dict[str, str] = {}
    if not path.exists():
        return values
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        values[key.strip()] = value.strip().strip('"').strip("'")
    return values


class RestError(RuntimeError):
    def __init__(self, message: str, status: int | None = None, body: str = "") -> None:
        super().__init__(message)
        self.status, self.body = status, body


class SupabaseRest:
    def __init__(self, url: str, key: str) -> None:
        self.url, self.key = url.rstrip("/"), key

    def request(self, method: str, table_path: str, query: list[tuple[str, str]] | None = None, payload: Any = None, prefer: str | None = None) -> Any:
        query_string = urllib.parse.urlencode(query or [])
        url = f"{self.url}/rest/v1/{table_path}" + (f"?{query_string}" if query_string else "")
        body = None if payload is None else json.dumps(payload, ensure_ascii=False).encode("utf-8")
        headers = {"apikey": self.key, "Authorization": f"Bearer {self.key}", "Content-Type": "application/json"}
        if prefer:
            headers["Prefer"] = prefer
        request = urllib.request.Request(url, data=body, headers=headers, method=method)
        try:
            with urllib.request.urlopen(request, timeout=20) as response:
                raw = response.read().decode("utf-8")
                return json.loads(raw) if raw else None
        except urllib.error.HTTPError as exc:
            body_text = exc.read().decode("utf-8", errors="replace")
            raise RestError(f"Supabase {method} {table_path} returned HTTP {exc.code}", exc.code, body_text) from exc
        except urllib.error.URLError as exc:
            raise RestError(f"Supabase connection failed: {exc.reason}") from exc


def load_database_context(env_path: Path) -> tuple[SupabaseRest | None, dict[str, Any]]:
    env = load_env(env_path)
    url, key = env.get("VITE_SUPABASE_URL", ""), env.get("VITE_SUPABASE_ANON_KEY", "")
    if not url or not key:
        return None, {"status": "UNAVAILABLE", "reason": "VITE_SUPABASE_URL/ANON_KEY tidak lengkap."}
    client = SupabaseRest(url, key)
    try:
        anggaran = client.request("GET", "anggaran_tahun", [("select", "id,tahun,nama_paket,lokasi,kecamatan,pagu_fisik,pagu_perencanaan,pagu_pengawasan,pagu_honor,total_pagu,sub_kegiatan_id,keterangan"), ("limit", "5000")])
        sub_kegiatan = client.request("GET", "sub_kegiatan", [("select", "id,kode,nama,aktif"), ("limit", "1000")])
        try:
            links = client.request("GET", "anggaran_ruas_link", [("select", "id,anggaran_id,spasial_ref,spasial_nama,link_type,sumber_link,confidence_score,catatan,spasial_type"), ("limit", "10000")])
        except RestError:
            links = client.request("GET", "anggaran_ruas_link", [("select", "id,anggaran_id,spasial_ref,spasial_nama,link_type,sumber_link,confidence_score,catatan"), ("limit", "10000")])
        return client, {"status": "AVAILABLE", "anggaran": anggaran or [], "sub_kegiatan": sub_kegiatan or [], "links": links or []}
    except RestError as exc:
        return None, {"status": "UNAVAILABLE", "reason": str(exc), "detail": exc.body[:500]}


def target_budget_field(record: dict[str, Any]) -> str:
    return "pagu_perencanaan" if record["jenis"] == "PLANNING" else "pagu_fisik"


def package_candidates(record: dict[str, Any], database: dict[str, Any]) -> list[dict[str, Any]]:
    source_name, source_sub_name = normalize_package_name(record["nama_paket"]), record.get("sub_name", "")
    sub_by_id = {str(item.get("id")): item for item in database.get("sub_kegiatan", [])}
    result = []
    for item in database.get("anggaran", []):
        if int(item.get("tahun") or 0) != record["target_year"] or normalize_package_name(item.get("nama_paket")) != source_name:
            continue
        item_sub = sub_by_id.get(str(item.get("sub_kegiatan_id")), {})
        if source_sub_name and normalize_sub_name(item_sub.get("nama")) != source_sub_name:
            continue
        result.append(item)
    return result


def resolve_package(record: dict[str, Any], database: dict[str, Any]) -> dict[str, Any]:
    record["target_anggaran_id"], record["database_status"] = "", database.get("status", "UNAVAILABLE")
    existing_id = record["existing_id"]
    if database.get("status") != "AVAILABLE":
        if existing_id:
            record["target_anggaran_id"] = existing_id
            record["package_action"] = "UPDATE_EXISTING" if record["tindakan"] == "UPDATE_EXISTING" else "EXISTING"
            record["notes"] = "Database belum dapat diaudit; Existing ID dipertahankan sebagai kandidat, bukan bukti row tersedia."
        else:
            record["package_action"] = "NEW"
            record["notes"] = "Database belum dapat diaudit; kandidat NEW belum boleh dieksekusi sebelum lookup online."
        return record

    by_id = {str(item.get("id")): item for item in database.get("anggaran", [])}
    if existing_id:
        if existing_id not in by_id:
            fallback = package_candidates(record, database)
            if len(fallback) == 1:
                item = fallback[0]
                amount = max(
                    float(item.get("pagu_fisik") or 0),
                    float(item.get("pagu_perencanaan") or 0),
                )
                if abs(amount - record["effective_pagu"]) <= 1:
                    record["target_anggaran_id"] = str(item.get("id"))
                    record["package_action"] = "UPDATE_EXISTING" if record["tindakan"] == "UPDATE_EXISTING" else "EXISTING"
                    record["notes"] = "Existing ID workbook tidak ditemukan, tetapi satu paket existing cocok berdasarkan tahun, nama, sub kegiatan, dan pagu."
                    return record
            record["package_action"], record["notes"] = "REVIEW", "Existing ID dari workbook tidak ditemukan dan tidak ada satu kandidat pengganti yang aman."
            return record
        record["target_anggaran_id"] = existing_id
        record["package_action"] = "UPDATE_EXISTING" if record["tindakan"] == "UPDATE_EXISTING" else "EXISTING"
        return record

    candidates = package_candidates(record, database)
    if len(candidates) == 1:
        amount = max(float(candidates[0].get("pagu_fisik") or 0), float(candidates[0].get("pagu_perencanaan") or 0))
        if abs(amount - record["effective_pagu"]) <= 1:
            record["target_anggaran_id"], record["package_action"] = str(candidates[0].get("id")), "EXISTING"
            record["notes"] = "Kandidat existing cocok pada tahun, nama, sub kegiatan, dan pagu efektif."
            return record
        record["package_action"], record["notes"] = "REVIEW", "Nama dan sub kegiatan cocok, tetapi pagu efektif berbeda; tidak membuat duplicate otomatis."
        return record
    if len(candidates) > 1:
        record["package_action"], record["notes"] = "REVIEW", "Lebih dari satu paket existing memiliki nama/sub kegiatan yang sama."
        return record
    if not record.get("sub_name"):
        record["package_action"], record["notes"] = "REVIEW", "Kode sub kegiatan tidak dapat diambil dari workbook."
        return record
    sub_matches = [item for item in database.get("sub_kegiatan", []) if normalize_sub_name(item.get("nama")) == record["sub_name"]]
    if len(sub_matches) != 1:
        record["package_action"], record["notes"] = "REVIEW", f"Sub kegiatan {record['sub_name']} tidak memiliki tepat satu master aktif."
        return record
    record["package_action"] = "NEW"
    return record


def add_road_fields(record: dict[str, Any], roads: list[dict[str, Any]]) -> dict[str, Any]:
    match = classify_road_match(record["nama_paket"], roads)
    record["road_match_status"], record["road_candidates"], record["road_notes"] = match["status"], match["candidates"], match["notes"]
    best = match["candidates"][0] if match["candidates"] else {}
    second = match["candidates"][1] if len(match["candidates"]) > 1 else {}
    if match["status"] in {"EXACT", "FUZZY_AUTO"}:
        record["matched_spasial_ref"], record["matched_spasial_nama"] = best.get("road_ref", ""), best.get("road_name", "")
        record["match_score"], record["match_method"] = best.get("score", ""), best.get("method", "")
    else:
        record["matched_spasial_ref"], record["matched_spasial_nama"], record["match_score"], record["match_method"] = "", "", "", ""
    record["second_best_score"] = second.get("score", "")
    return record


def build_rows(records: list[dict[str, Any]], roads: list[dict[str, Any]], database: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for source in records:
        record = dict(source)
        record["notes"] = record.get("catatan_source", "")
        if record["status"] in {"CANCELLED", "REPLACED"} or record["tindakan"] in {"CANCEL", "REPLACE"}:
            record.update({
                "package_action": "SKIPPED", "road_match_status": "SKIPPED_CANCELLED", "target_anggaran_id": "",
                "matched_spasial_ref": "", "matched_spasial_nama": "", "match_score": "", "match_method": "",
                "second_best_score": "", "road_candidates": [], "database_status": database.get("status", "UNAVAILABLE"),
                "road_notes": "Status/tindakan cancelled/replaced; tidak diimport dan tidak dilink.",
            })
        else:
            add_road_fields(record, roads)
            resolve_package(record, database)
            if record.get("road_notes"):
                record["notes"] = "; ".join(filter(None, [record.get("notes"), record["road_notes"]]))
        rows.append(record)
    return rows


def csv_value(value: Any) -> Any:
    return int(value) if isinstance(value, float) and value.is_integer() else ("" if value is None else value)


def verify_after_write(rows: list[dict[str, Any]], fresh_database: dict[str, Any], pre_write_database: dict[str, Any]) -> dict[str, Any]:
    target_rows = [row for row in rows if row.get("package_action") in {"NEW", "EXISTING", "UPDATE_EXISTING"} and row.get("road_match_status") != "SKIPPED_CANCELLED"]
    fresh_ids = {str(item.get("id")) for item in fresh_database.get("anggaran", [])}
    missing_target_ids = [row.get("source_row") for row in target_rows if row.get("target_anggaran_id") and str(row.get("target_anggaran_id")) not in fresh_ids]
    link_keys = [
        (str(link.get("anggaran_id")), str(link.get("spasial_ref")), str(link.get("link_type")))
        for link in fresh_database.get("links", [])
    ]
    duplicate_link_keys = sum(count - 1 for count in Counter(link_keys).values() if count > 1)
    manual_before = sum(str(link.get("sumber_link", "")).lower() == "manual" for link in pre_write_database.get("links", []))
    manual_after = sum(str(link.get("sumber_link", "")).lower() == "manual" for link in fresh_database.get("links", []))
    duplicate_source_keys = [
        key for key, count in Counter((row.get("target_year"), normalize_package_name(row.get("nama_paket")), row.get("sub_name", "")) for row in target_rows).items()
        if count > 1
    ]
    return {
        "database_status": fresh_database.get("status"),
        "target_rows_missing_after_write": missing_target_ids,
        "duplicate_spatial_link_rows": duplicate_link_keys,
        "duplicate_source_package_keys_in_safe_batch": [list(key) for key in duplicate_source_keys],
        "manual_links_before": manual_before,
        "manual_links_after": manual_after,
        "manual_links_lost": max(0, manual_before - manual_after),
        "target_year_counts": dict(Counter(int(item.get("tahun") or 0) for item in fresh_database.get("anggaran", []))),
        "cancelled_replaced_rows_written": 0,
        "passed": not missing_target_ids and duplicate_link_keys == 0 and not duplicate_source_keys and manual_before == manual_after,
    }


def verify_current_state(rows: list[dict[str, Any]], database: dict[str, Any]) -> dict[str, Any]:
    safe_rows = [
        row for row in rows
        if row.get("package_action") in {"EXISTING", "NEW", "UPDATE_EXISTING"}
        and row.get("road_match_status") != "SKIPPED_CANCELLED"
    ]
    current_packages = {str(item.get("id")): item for item in database.get("anggaran", [])}
    expected_ids = {str(row.get("target_anggaran_id")) for row in safe_rows if row.get("target_anggaran_id")}
    missing_ids = sorted(expected_ids - set(current_packages))
    expected_links = {
        (str(row.get("target_anggaran_id")), str(row.get("matched_spasial_ref")), "perencanaan" if row.get("jenis") == "PLANNING" else "fisik")
        for row in safe_rows
        if row.get("road_match_status") in {"EXACT", "FUZZY_AUTO"} and row.get("matched_spasial_ref")
    }
    actual_links = {
        (str(link.get("anggaran_id")), str(link.get("spasial_ref")), str(link.get("link_type")))
        for link in database.get("links", [])
        if link.get("sumber_link") == "auto_rekonsiliasi_apbdp_2026"
    }
    all_link_keys = [
        (str(link.get("anggaran_id")), str(link.get("spasial_ref")), str(link.get("link_type")))
        for link in database.get("links", [])
    ]
    duplicate_link_rows = sum(count - 1 for count in Counter(all_link_keys).values() if count > 1)
    cancelled_ids = {row.get("existing_id") for row in rows if row.get("road_match_status") == "SKIPPED_CANCELLED" and row.get("existing_id")}
    safe_source_ids = [str(row.get("target_anggaran_id")) for row in safe_rows]
    source_id_duplicates = sum(count - 1 for count in Counter(safe_source_ids).values() if count > 1)
    manual_links = [link for link in database.get("links", []) if str(link.get("sumber_link", "")).lower() == "manual"]
    return {
        "database_status": database.get("status"),
        "package_count": len(current_packages),
        "target_year_counts": dict(Counter(int(item.get("tahun") or 0) for item in current_packages.values())),
        "safe_source_rows": len(safe_rows),
        "safe_source_ids_missing": missing_ids,
        "safe_source_id_duplicates": source_id_duplicates,
        "cancelled_replaced_ids_present": sorted(cancelled_ids.intersection(current_packages)),
        "expected_auto_links": len(expected_links),
        "actual_auto_links": len(actual_links),
        "auto_links_missing": [list(key) for key in sorted(expected_links - actual_links)],
        "auto_links_unexpected": [list(key) for key in sorted(actual_links - expected_links)],
        "all_spatial_link_duplicate_rows": duplicate_link_rows,
        "manual_links_current": len(manual_links),
        "review_rows_untouched": sum(row.get("package_action") == "REVIEW" for row in rows),
        "cancelled_replaced_rows_skipped": sum(row.get("road_match_status") == "SKIPPED_CANCELLED" for row in rows),
        "passed": (
            database.get("status") == "AVAILABLE"
            and not missing_ids
            and source_id_duplicates == 0
            and not cancelled_ids.intersection(current_packages)
            and expected_links == actual_links
            and duplicate_link_rows == 0
            and len(manual_links) == 61
        ),
    }


def write_reports(rows: list[dict[str, Any]], output_dir: Path, workbook_summary: dict[str, Any], database: dict[str, Any], roads_count: int, writes: dict[str, Any] | None = None, verification: dict[str, Any] | None = None) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    with (output_dir / "rekonsiliasi_apbdp_2026_spatial_dry_run.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=REPORT_COLUMNS, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({column: csv_value(row.get(column, "")) for column in REPORT_COLUMNS})

    with (output_dir / "rekonsiliasi_apbdp_2026_manual_review.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=MANUAL_COLUMNS)
        writer.writeheader()
        for row in rows:
            if row.get("package_action") != "REVIEW" and row.get("road_match_status") not in {"REVIEW", "UNMATCHED"}:
                continue
            candidates = [{"ref": item.get("road_ref"), "name": item.get("road_name"), "score": item.get("score"), "method": item.get("method")} for item in row.get("road_candidates", [])]
            writer.writerow({
                "source_sheet": row.get("source_sheet"), "source_row": row.get("source_row"), "target_year": row.get("target_year"),
                "jenis": row.get("jenis"), "nama_paket_source": row.get("nama_paket_source"), "sub_source": row.get("sub_source"),
                "effective_pagu": csv_value(row.get("effective_pagu")), "package_action": row.get("package_action"),
                "road_match_status": row.get("road_match_status"), "notes": row.get("notes"),
                "top_candidates": json.dumps(candidates, ensure_ascii=False),
            })

    summary = {
        "workbook": str(DEFAULT_WORKBOOK), "sheets_inspected": workbook_summary.get("sheet_names", []),
        "source_main_rows": workbook_summary.get("main_rows", 0), "action_sheet_summary": workbook_summary.get("action_counts", {}),
        "main_jenis_counts": workbook_summary.get("all_main_jenis", {}), "main_status_counts": workbook_summary.get("all_main_status", {}),
        "main_tindakan_counts": workbook_summary.get("all_main_tindakan", {}),
        "eligible_rows": {
            "planning": sum(row.get("jenis") == "PLANNING" and row.get("road_match_status") != "SKIPPED_CANCELLED" for row in rows),
            "physical": sum(row.get("jenis") == "PHYSICAL" and row.get("road_match_status") != "SKIPPED_CANCELLED" for row in rows),
        },
        "skipped_cancelled_replaced": sum(row.get("road_match_status") == "SKIPPED_CANCELLED" for row in rows),
        "roads_loaded": roads_count, "package_actions": dict(Counter(row.get("package_action") for row in rows)),
        "road_match_statuses": dict(Counter(row.get("road_match_status") for row in rows)),
        "database_status": database.get("status", "UNAVAILABLE"), "database_reason": database.get("reason", ""),
        "writes": writes or {"performed": False, "packages": 0, "links": 0, "manual_links_preserved": 0},
        "verification": {
            "cancelled_replaced_inserted": 0,
            "duplicate_creation_check": "NOT_RUN_OFFLINE" if database.get("status") != "AVAILABLE" else "PENDING_EXECUTION",
            "manual_link_check": "PENDING_EXECUTION" if database.get("status") == "AVAILABLE" else "NOT_RUN_OFFLINE",
        },
    }
    if verification is not None:
        summary["post_write_verification"] = verification
        (output_dir / "rekonsiliasi_apbdp_2026_post_write_verification.json").write_text(json.dumps(verification, ensure_ascii=False, indent=2), encoding="utf-8")
    (output_dir / "rekonsiliasi_apbdp_2026_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")


def package_payload(record: dict[str, Any], sub_id: str | None) -> dict[str, Any]:
    payload = {
        "tahun": record["target_year"], "nama_paket": record["nama_paket"],
        "pagu_fisik": record["effective_pagu"] if record["jenis"] == "PHYSICAL" else 0,
        "pagu_perencanaan": record["effective_pagu"] if record["jenis"] == "PLANNING" else 0,
        "pagu_pengawasan": 0, "pagu_honor": 0,
        "keterangan": f"Sumber: REKONSILIASI_APBDP_2026.xlsx; sheet={record['source_sheet']}; row={record['source_row']}; status={record['status']}",
    }
    if sub_id:
        payload["sub_kegiatan_id"] = sub_id
    return payload


def execute_safe(rows: list[dict[str, Any]], database: dict[str, Any], client: SupabaseRest) -> dict[str, Any]:
    if database.get("status") != "AVAILABLE":
        raise RuntimeError("Database tidak tersedia; execute dibatalkan tanpa write.")
    sub_by_name = defaultdict(list)
    for item in database.get("sub_kegiatan", []):
        sub_by_name[normalize_sub_name(item.get("nama"))].append(item)
    links_by_package: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for link in database.get("links", []):
        links_by_package[str(link.get("anggaran_id"))].append(link)

    created_keys: dict[tuple[int, str, str], tuple[str, float]] = {}
    package_writes = link_writes = manual_preserved = skipped = 0
    for row in rows:
        if row.get("road_match_status") == "SKIPPED_CANCELLED" or row.get("package_action") not in {"NEW", "EXISTING", "UPDATE_EXISTING"}:
            skipped += 1
            continue
        code = row.get("sub_code", "")
        sub_matches = sub_by_name.get(row.get("sub_name", ""), []) if row.get("sub_name") else []
        sub_id = str(sub_matches[0].get("id")) if len(sub_matches) == 1 else None
        if row.get("package_action") == "NEW" and not sub_id:
            row["package_action"], row["notes"] = "REVIEW", "; ".join(filter(None, [row.get("notes"), "Master sub kegiatan tidak unik."]))
            skipped += 1
            continue

        package_key = (int(row["target_year"]), normalize_text(row["nama_paket"]), code)
        target_id = row.get("target_anggaran_id")
        if row.get("package_action") == "NEW":
            if package_key in created_keys:
                previous_id, previous_amount = created_keys[package_key]
                if abs(previous_amount - row["effective_pagu"]) <= 1:
                    target_id, row["package_action"] = previous_id, "EXISTING"
                else:
                    row["package_action"], row["notes"] = "REVIEW", "; ".join(filter(None, [row.get("notes"), "Duplicate nama/sub kegiatan dengan pagu berbeda dalam batch; perlu review."]))
                    skipped += 1
                    continue
            else:
                response = client.request("POST", "anggaran_tahun", payload=package_payload(row, sub_id), prefer="return=representation")
                inserted = response[0] if isinstance(response, list) and response else response
                target_id = str(inserted.get("id")) if isinstance(inserted, dict) else ""
                if not target_id:
                    raise RuntimeError(f"Insert anggaran_tahun tidak mengembalikan id untuk row {row['source_row']}.")
                created_keys[package_key] = (target_id, row["effective_pagu"])
                package_writes += 1
        if not target_id:
            skipped += 1
            continue
        row["target_anggaran_id"] = target_id
        if row.get("package_action") in {"EXISTING", "UPDATE_EXISTING"}:
            patch_payload = {
                "tahun": row["target_year"],
                "pagu_fisik": row["effective_pagu"] if row["jenis"] == "PHYSICAL" else 0,
                "pagu_perencanaan": row["effective_pagu"] if row["jenis"] == "PLANNING" else 0,
            }
            if sub_id:
                patch_payload["sub_kegiatan_id"] = sub_id
            client.request("PATCH", "anggaran_tahun", query=[("id", f"eq.{target_id}")], payload=patch_payload, prefer="return=minimal")
            package_writes += 1
        if row.get("road_match_status") not in {"EXACT", "FUZZY_AUTO"}:
            continue

        current_links = links_by_package.get(str(target_id), [])
        if any(str(link.get("sumber_link", "")).lower() == "manual" for link in current_links):
            manual_preserved += 1
            row["notes"] = "; ".join(filter(None, [row.get("notes"), "Manual spatial link dipertahankan; auto-link tidak menimpa."]))
            continue

        best = row["road_candidates"][0]
        payload = {
            "anggaran_id": target_id, "spasial_ref": best["road_ref"], "spasial_nama": best["road_name"],
            "link_type": "perencanaan" if row["jenis"] == "PLANNING" else "fisik",
            "sumber_link": "auto_rekonsiliasi_apbdp_2026", "confidence_score": best["score"],
            "catatan": f"{best['method']}: {best['reason']}; source row {row['source_row']}", "spasial_type": "road",
        }
        try:
            client.request("POST", "anggaran_ruas_link?on_conflict=anggaran_id%2Cspasial_ref%2Clink_type", payload=payload, prefer="resolution=merge-duplicates,return=representation")
        except RestError as exc:
            if "spasial_type" not in exc.body.lower() and "column" not in exc.body.lower():
                raise
            payload.pop("spasial_type", None)
            client.request("POST", "anggaran_ruas_link?on_conflict=anggaran_id%2Cspasial_ref%2Clink_type", payload=payload, prefer="resolution=merge-duplicates,return=representation")
        link_writes += 1

    return {"performed": True, "packages": package_writes, "links": link_writes, "manual_links_preserved": manual_preserved, "skipped": skipped}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--roads", type=Path, default=DEFAULT_ROADS)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--env-file", type=Path, default=REPO_ROOT / ".env")
    parser.add_argument("--execute", action="store_true", help="Eksekusi package/link writes yang lolos filter.")
    parser.add_argument("--approve-dry-run", action="store_true", help="Wajib bersama --execute setelah review dry-run.")
    parser.add_argument("--verify-only", action="store_true", help="Verifikasi read-only terhadap state database saat ini.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if not args.workbook.exists():
        print(f"Workbook tidak ditemukan: {args.workbook}", file=sys.stderr)
        return 2
    if not args.roads.exists():
        print(f"Sumber jalan tidak ditemukan: {args.roads}", file=sys.stderr)
        return 2
    if args.execute and not args.approve_dry_run:
        print("Refuse execute: gunakan --execute --approve-dry-run setelah dry-run direview.", file=sys.stderr)
        return 2
    if args.execute and args.verify_only:
        print("Pilih salah satu: --execute atau --verify-only.", file=sys.stderr)
        return 2

    records, workbook_summary = read_workbook(args.workbook)
    roads = load_road_source(args.roads)
    client, database = load_database_context(args.env_file)
    rows = build_rows(records, roads, database)
    writes = {"performed": False, "packages": 0, "links": 0, "manual_links_preserved": 0}
    verification = None

    if args.verify_only:
        verification = verify_current_state(rows, database)
        write_reports(rows, args.output_dir, workbook_summary, database, len(roads), writes, verification)
        print(json.dumps(verification, ensure_ascii=False, indent=2))
        return 0 if verification.get("passed") else 3

    if args.execute:
        if client is None or database.get("status") != "AVAILABLE":
            print("Database tidak tersedia; tidak ada write yang dilakukan.", file=sys.stderr)
            write_reports(rows, args.output_dir, workbook_summary, database, len(roads), writes)
            return 2
        writes = execute_safe(rows, database, client)
        _, fresh_database = load_database_context(args.env_file)
        verification = verify_after_write(rows, fresh_database, database)
        database = fresh_database
        if not verification.get("passed"):
            print("Post-write verification failed; inspect summary JSON.", file=sys.stderr)
            write_reports(rows, args.output_dir, workbook_summary, fresh_database, len(roads), writes, verification)
            return 3

    write_reports(rows, args.output_dir, workbook_summary, database, len(roads), writes, verification)
    print(json.dumps({
        "database_status": database.get("status"),
        "planning": sum(row["jenis"] == "PLANNING" and row["road_match_status"] != "SKIPPED_CANCELLED" for row in rows),
        "physical": sum(row["jenis"] == "PHYSICAL" and row["road_match_status"] != "SKIPPED_CANCELLED" for row in rows),
        "package_actions": dict(Counter(row.get("package_action") for row in rows)),
        "road_match_statuses": dict(Counter(row.get("road_match_status") for row in rows)),
        "reports": [str(args.output_dir / name) for name in ["rekonsiliasi_apbdp_2026_spatial_dry_run.csv", "rekonsiliasi_apbdp_2026_manual_review.csv", "rekonsiliasi_apbdp_2026_summary.json"]],
        "writes": writes,
    }, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
